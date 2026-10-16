"""천사의시_대사_번역.xlsx (+ 천사의시_3줄초과_수정필요.xlsx가 있으면 병합) -> tenshi_master.csv

사용법: python xlsx_to_csv.py <translation 폴더> <출력 csv>
"""
import csv
import os
import sys

from openpyxl import load_workbook

base = sys.argv[1]
out_csv = sys.argv[2]

MASTER = os.path.join(base, "천사의시_대사_번역.xlsx")
OVER3 = os.path.join(base, "천사의시_3줄초과_수정필요.xlsx")

COLS = ["offset", "key_hex", "jp", "en", "ko"]

def read_xlsx(path):
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = []
    header = None
    for r in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c) if c is not None else "" for c in r]
            continue
        d = {}
        for i, name in enumerate(header):
            if name in COLS:
                v = r[i] if i < len(r) else None
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                d[name] = "" if v is None else str(v).replace("_x000D_", "\r")
        rows.append(d)
    wb.close()
    return rows

master = read_xlsx(MASTER)
print(f"master rows: {len(master)}")

if os.path.exists(OVER3):
    over = read_xlsx(OVER3)
    by_offset = {r["offset"]: r for r in master}
    merged = 0
    for r in over:
        m = by_offset.get(r["offset"])
        if m is not None and r.get("ko", "") != m.get("ko", ""):
            m["ko"] = r["ko"]
            merged += 1
    print(f"merged from over3 file: {merged}")

with open(out_csv, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=COLS)
    w.writeheader()
    for r in master:
        w.writerow({c: r.get(c, "") for c in COLS})
print(f"written: {out_csv}")
