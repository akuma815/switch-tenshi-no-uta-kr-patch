import csv
import sys
import os

sys.path.insert(0, os.path.dirname(__file__))
from rewrap_ko import wrap_text

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def csv_to_xlsx(csv_path, xlsx_path, col_widths, add_linecount=False):
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    fieldnames = list(rows[0].keys())
    if add_linecount and "줄수" not in fieldnames:
        fieldnames = fieldnames + ["줄수(자동줄바꿈 후)"]

    wb = Workbook()
    ws = wb.active
    ws.title = "번역"

    for ci, name in enumerate(fieldnames, 1):
        c = ws.cell(row=1, column=ci, value=name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    for ri, row in enumerate(rows, 2):
        for ci, name in enumerate(fieldnames, 1):
            if name == "줄수(자동줄바꿈 후)":
                ko = (row.get("ko") or "").strip()
                val = wrap_text(ko).count("\n") + 1 if ko else 0
            else:
                val = row.get(name, "")
            c = ws.cell(row=ri, column=ci, value=val)
            c.alignment = WRAP

    for ci, name in enumerate(fieldnames, 1):
        letter = get_column_letter(ci)
        ws.column_dimensions[letter].width = col_widths.get(name, 12)

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"{os.path.basename(xlsx_path)}: {len(rows)} rows")

BASE = r"G:\Downloads\Tenshi no Uta\translation"

csv_to_xlsx(
    os.path.join(BASE, "tenshi_text_en_paired.csv"),
    os.path.join(BASE, "천사의시_대사_번역.xlsx"),
    {"offset": 9, "key_hex": 14, "jp": 45, "en": 45, "ko": 50},
    add_linecount=True,
)
csv_to_xlsx(
    os.path.join(BASE, "tenshi_over3lines.csv"),
    os.path.join(BASE, "천사의시_3줄초과_수정필요.xlsx"),
    {"offset": 9, "key_hex": 14, "jp": 45, "en": 45, "ko": 50},
    add_linecount=True,
)
