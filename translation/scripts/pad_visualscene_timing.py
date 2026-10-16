# -*- coding: utf-8 -*-
import csv
from pathlib import Path

from openpyxl import load_workbook


PAD = "\u3000"
OLD_PAD = "\u200b"


def clean_padding(text: str) -> str:
    return (text or "").replace(OLD_PAD, "").rstrip(" ").rstrip(PAD)


def count_for_timing(text: str) -> int:
    text = clean_padding(text).replace("\\n", "\n").replace("\r\n", "\n").replace("\r", "\n")
    # Count line breaks too. The visual scene player appears to advance by text length,
    # so keep the count close to the shipped English strings.
    return len(text)


def padded_ko(en: str, ko: str) -> str:
    visible = clean_padding(ko)
    diff = count_for_timing(en) - count_for_timing(visible)
    if diff <= 0:
        return visible
    return visible + (PAD * diff)


def update_csv(path: Path) -> int:
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
        fieldnames = list(rows[0].keys()) if rows else []
    changed = 0
    for row in rows:
        ko = row.get("ko") or ""
        if not ko:
            continue
        new_ko = padded_ko(row.get("en") or "", ko)
        if new_ko != ko:
            row["ko"] = new_ko
            changed += 1
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return changed


def update_xlsx(path: Path) -> int:
    if not path.exists():
        return 0
    wb = load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    en_col = header.index("en") + 1
    ko_col = header.index("ko") + 1
    changed = 0
    for row_idx in range(2, ws.max_row + 1):
        en = ws.cell(row_idx, en_col).value or ""
        ko = ws.cell(row_idx, ko_col).value or ""
        if not ko:
            continue
        new_ko = padded_ko(str(en), str(ko))
        if new_ko != ko:
            ws.cell(row_idx, ko_col).value = new_ko
            changed += 1
    try:
        wb.save(path)
    finally:
        wb.close()
    return changed


def report_csv(path: Path) -> None:
    if not path.exists():
        return
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    padded = 0
    total_pad = 0
    still_short = 0
    for row in rows:
        ko = row.get("ko") or ""
        pad_count = ko.count(PAD)
        if pad_count:
            padded += 1
            total_pad += pad_count
        if row.get("ko") and count_for_timing(row.get("ko") or "") < count_for_timing(row.get("en") or ""):
            still_short += 1
    print(f"padded_rows={padded}, total_fullwidth_space_padding={total_pad}, still_short={still_short}")


def clean_xlsx(path: Path) -> int:
    if not path.exists():
        return 0
    wb = load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    ko_col = header.index("ko") + 1
    changed = 0
    for row_idx in range(2, ws.max_row + 1):
        ko = ws.cell(row_idx, ko_col).value
        if not isinstance(ko, str):
            continue
        cleaned = clean_padding(ko)
        if cleaned != ko:
            ws.cell(row_idx, ko_col).value = cleaned
            changed += 1
    try:
        wb.save(path)
    finally:
        wb.close()
    return changed


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    csv_path = root / "work" / "visualscene_paired.csv"
    xlsx_path = root / "visualscene_paired.xlsx"
    try:
        print(f"cleaned xlsx padding: {clean_xlsx(xlsx_path)}")
    except PermissionError as exc:
        print(f"xlsx locked, skipped cleaning: {exc}")
    print(f"updated work csv: {update_csv(csv_path)}")
    report_csv(csv_path)


if __name__ == "__main__":
    main()
