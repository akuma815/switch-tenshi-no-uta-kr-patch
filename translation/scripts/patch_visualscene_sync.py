# -*- coding: utf-8 -*-
import csv
from pathlib import Path

from openpyxl import load_workbook


UPDATES = {
    "7429": "패왕 루키펠이 부활했다면,\n천상계 왕가의 피만이 그를 막는다",
    "7815": "우리 천상계와 지하계의 백성은\n인간을 위협해서는 안 됩니다",
    "8421": "옛날 당신의 아버지는 루키펠과 싸워\n이 보석의 힘으로 분쟁을 끝냈습니다.",
    "8663": "이제 당신이 아버지를 대신해\n루키펠을 물리쳐야 합니다",
    "9067": "클레어, 이제 로크란으로 가야 합니다.",
    "10977": "지하 길은 막혔다. 세상은 구원받았다.",
    "11460": "청년아, 슬퍼하지 마십시오.\n그 아이는 사명을 마치고 하늘로 돌아갑니다.",
    "17240": "자, 제가 누구일까요?\n마리아 역의 타카시마 마사라입니다.",
    "17671": "여기까지 오다니 대단하구나.\n정말 훌륭한 플레이어야!",
    "17792": "그럼 보상으로 한마디.\n“샹샹샹샹! 축하해!”",
    "17968": "지토 역의 타나카 히데유키입니다.",
    "18251": "지토 대사가 적어 아쉬웠지만,\nRPG를 좋아해서\n이 게임에 전력을 다했습니다.",
    "18456": "『천사의 시』 끝까지 보느라 수고 많으셨어요.\n클레어 역의 에모리 히로코입니다.",
    "18782": "재미있었어?\n이제 공부도 열심히 해야지. 안녕!",
    "19062": "엔야는 귀엽고 활기찬 소녀죠.\n그 매력이 잘 전해졌길 바랍니다.\n게임 클리어 축하드려요!",
    "19452": "안녕하세요. 제 목소리 아시겠나요?\n부젠 역의 미야우치 코헤이입니다.\n성우라 얼굴보다 이름이 익숙하실 겁니다.\n만나서 반갑습니다.",
    "20324": "하이디의 할아버지, 잇큐상의 스님.\n이 목소리, 아시겠죠?\n성우 일이 참 재미있습니다.\n꿈이 있다면 열심히 공부하세요.\n그럼 이만.",
    "20782": "여보세요. 테스트, 테스트.\n오늘은 날씨가 맑습니다.\n캐릭터 디자인의 후지 히로시입니다.",
    "21568": "『천사의 시』 재미있게 즐기셨나요?\n반응이 좋으면 후속작도 생각 중입니다.\n감상도 많이 보내 주세요.",
    "21961": "옛날, 평화로운 마을에 한 부부가 살았습니다.\n두 사람은 평범하게 사랑하고 결혼했습니다.\n다만 아내가 천사라는 점만 달랐습니다.",
    "22653": "으으... 어젯밤도 늦게 잤어.\n조금만 더 자게 해 줘.",
    "23267": "아, 클레어. 화내지 마. 일어날게.",
    "23945": "제대로 벌지도 못하면서\n그런 말은 잘도 하는군.",
    "24558": "젠장, 결혼 전엔 그렇게 귀여웠는데.\n이젠 돈 타령만 하며 살만 쪘잖아.\n왜 하필 나야!",
    "24886": "흥! 으악!",
    "25032": "괭이가 부러졌어...\n아아, 어떡하지. 또 클레어한테 맞겠네.",
    "25217": "그 무렵, 평화로운 세상에서\n검사의 길이 막힌 지토는...",
    "25579": "에호, 에호...\n이 사람아, 일 좀 하세요.",
    "26131": "으으... 난... 힉...\n검을 위해 사는 것 말고는 길이 없어!",
    "26305": "하지만 당신! 아이들을 좀 봐요!\n저렇게 수척해졌는데... 에호.",
    "26566": "원망하려면...\n이 평화로운 시대를 원망해라! 힉...",
    "27148": "미안해... 좀 참아 줘... 에헴...\n아빠가 일만 해 주면 말이지, 얘야.",
    "27741": "뭐 그렇게 멍하니 서 있어?\n어서 밥 준비나 해.",
    "28762": "으악! 죄송해요!",
    "29175": "그렇게 평화로운 나날은\n몇 년이고 계속되었다고 한다.",
}


def update_csv(path: Path) -> int:
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
        fieldnames = list(rows[0].keys()) if rows else []
    changed = 0
    for row in rows:
        value = UPDATES.get(str(row.get("offset", "")))
        if value is not None:
            row["ko"] = value
            changed += 1
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return changed


def update_xlsx(path: Path) -> int:
    if not path.exists():
        return 0
    wb = load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    offset_col = header.index("offset") + 1
    ko_col = header.index("ko") + 1
    changed = 0
    for row_idx in range(2, ws.max_row + 1):
        offset = str(ws.cell(row_idx, offset_col).value or "")
        value = UPDATES.get(offset)
        if value is not None:
            ws.cell(row_idx, ko_col).value = value
            changed += 1
    try:
        wb.save(path)
    finally:
        wb.close()
    return changed


def char_width(text: str) -> float:
    return sum(0.5 if (ord(ch) < 0x2000 or 0xFF61 <= ord(ch) <= 0xFF9F) else 1 for ch in text)


def verify_csv(path: Path) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    line_more = []
    long_lines = []
    for row in rows:
        en_lines = (row.get("en") or "").replace("\\n", "\n").split("\n")
        ko_lines = (row.get("ko") or "").replace("\\n", "\n").split("\n")
        max_width = max(char_width(line) for line in ko_lines) if ko_lines else 0
        if len(ko_lines) > len(en_lines):
            line_more.append((row.get("offset"), len(en_lines), len(ko_lines)))
        if max_width > 24:
            long_lines.append((row.get("offset"), max_width))
    print(f"line_more={len(line_more)} {line_more[:20]}")
    print(f"long_lines_gt_24={len(long_lines)} {long_lines[:20]}")


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    csv_changed = update_csv(root / "work" / "visualscene_paired.csv")
    print(f"updated work csv: {csv_changed}")
    try:
        xlsx_changed = update_xlsx(root / "visualscene_paired.xlsx")
        print(f"updated xlsx: {xlsx_changed}")
    except PermissionError as exc:
        print(f"xlsx locked, skipped: {exc}")
    verify_csv(root / "work" / "visualscene_paired.csv")


if __name__ == "__main__":
    main()
